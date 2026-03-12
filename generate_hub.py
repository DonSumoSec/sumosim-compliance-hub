#!/usr/bin/env python3
"""
SumoSim Compliance Hub Generator v2
Reads SumoSim_Compliance_Hub.xlsx and generates a self-contained index.html
Run: python3 generate_hub.py
"""

import sys, os, json
from datetime import datetime, date, timedelta

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

XL_FILE  = "SumoSim_Compliance_Hub.xlsx"
OUT_FILE = "index.html"

def fmt_date(val):
    if val is None: return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%d %b %Y") if hasattr(val, 'hour') else val.strftime("%d %b %Y")
    v = str(val).strip()
    if v.replace('.','').isdigit() and float(v) > 40000:
        return (date(1899,12,30) + timedelta(days=int(float(v)))).strftime("%d %b %Y")
    return v

def load_data(path):
    if not os.path.exists(path):
        print(f"ERROR: {path} not found"); sys.exit(1)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        print(f"ERROR opening {path}: {e}"); sys.exit(1)

    dash  = next((s for s in wb.sheetnames if "Dashboard" in s), None)
    cont  = next((s for s in wb.sheetnames if "Content"   in s), None)
    if not dash:
        print(f"ERROR: No Dashboard sheet. Sheets: {wb.sheetnames}"); sys.exit(1)

    docs = []
    for row in wb[dash].iter_rows(min_row=7, max_row=60, values_only=True):
        if not row[1]: continue
        t = str(row[1]).strip()
        if not t or t.startswith("TOTAL") or t == "DOCUMENT": continue
        docs.append({
            "title":       t,
            "category":    str(row[2]).strip() if row[2] else "",
            "version":     str(row[3]).strip() if row[3] else "",
            "status":      str(row[4]).strip() if row[4] else "",
            "reviewedBy":  str(row[5]).strip() if row[5] else "",
            "lastReviewed": fmt_date(row[6]),
            "nextReview":   fmt_date(row[7]),
            "websitePage": str(row[8]).strip() if row[8] else "",
            "published":   str(row[9]).strip() if row[9] else "",
        })

    content = {}
    if cont:
        for row in wb[cont].iter_rows(min_row=5, max_row=300, values_only=True):
            if not row[1] or str(row[1]).strip() in ("DOCUMENT NAME",""):  continue
            t = str(row[1]).strip()
            if t not in content: content[t] = []
            content[t].append({
                "section": str(row[2]).strip() if row[2] else "",
                "live":    str(row[3]).strip() if row[3] and str(row[3]) not in ("—","None") else "",
                "draft":   str(row[4]).strip() if row[4] and str(row[4]) not in ("—","None") else "",
                "notes":   str(row[5]).strip() if row[5] and str(row[5]) != "None" else "",
            })

    return docs, content

def generate_html(docs, content):
    generated = datetime.now().strftime("%d %b %Y %H:%M")
    dj = json.dumps(docs,    ensure_ascii=False)
    cj = json.dumps(content, ensure_ascii=False)
    live  = sum(1 for d in docs if d["status"]=="Live")
    draft = sum(1 for d in docs if d["status"]=="Draft")
    pub   = sum(1 for d in docs if d["published"]=="Yes")

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1.0">
<title>SumoSim | Compliance Hub</title>
<style>
*{{margin:0;padding:0;box-sizing:border-box}}
body{{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;background:#0a0a0a;color:#e0e0e0;min-height:100vh}}
.topbar{{background:#111;border-bottom:1px solid #222;padding:12px 24px;display:flex;align-items:center;justify-content:space-between}}
.brand{{font-size:15px;font-weight:600}}.brand span{{color:#fff}}.brand .dot{{color:#666;margin:0 8px}}
.meta{{font-size:12px;color:#555}}
.layout{{display:flex;height:calc(100vh - 45px)}}
.sidebar{{width:230px;background:#111;border-right:1px solid #1e1e1e;overflow-y:auto;flex-shrink:0;padding:16px 0}}
.sb-sec{{padding:8px 16px 4px;font-size:10px;font-weight:700;color:#444;letter-spacing:1.5px;text-transform:uppercase}}
.sb-item{{padding:8px 16px;font-size:13px;color:#888;cursor:pointer;border-left:2px solid transparent;transition:all .15s;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}}
.sb-item:hover{{color:#ccc;background:#161616}}.sb-item.active{{color:#fff;border-left-color:#fff;background:#161616}}
.badge{{display:inline-block;font-size:9px;padding:1px 5px;border-radius:3px;margin-left:6px;font-weight:600;vertical-align:middle}}
.bl{{background:#1a3a1a;color:#4caf50}}.bd{{background:#3a2a1a;color:#ff9800}}.br{{background:#1a2a3a;color:#2196f3}}.bp{{background:#2a2a1a;color:#ffeb3b}}
.main{{flex:1;overflow-y:auto;padding:24px}}
.grid{{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin-bottom:28px}}
.card{{background:#111;border:1px solid #1e1e1e;border-radius:6px;padding:16px}}
.num{{font-size:28px;font-weight:700;color:#fff}}.lbl{{font-size:11px;color:#555;margin-top:4px;text-transform:uppercase;letter-spacing:.5px}}
table{{width:100%;border-collapse:collapse;background:#111;border-radius:6px;overflow:hidden;border:1px solid #1e1e1e}}
th{{background:#161616;padding:10px 14px;text-align:left;font-size:11px;color:#555;text-transform:uppercase;letter-spacing:.8px;font-weight:600;border-bottom:1px solid #1e1e1e}}
td{{padding:10px 14px;font-size:13px;border-bottom:1px solid #161616}}
tr:last-child td{{border-bottom:none}}tr:hover td{{background:#141414;cursor:pointer}}
.pill{{display:inline-block;font-size:10px;padding:2px 8px;border-radius:12px;font-weight:600}}
.sL{{background:#1a3a1a;color:#4caf50}}.sD{{background:#3a2a1a;color:#ff9800}}.sR{{background:#1a2a3a;color:#2196f3}}.sP{{background:#2a2a1a;color:#ffeb3b}}
.pY{{color:#4caf50;font-size:12px}}.pN{{color:#555;font-size:12px}}.pP{{color:#ff9800;font-size:12px}}
.panel{{background:#111;border:1px solid #1e1e1e;border-radius:6px;padding:24px}}
.ptitle{{font-size:20px;font-weight:700;color:#fff;margin-bottom:16px}}
.pmeta{{display:flex;gap:16px;flex-wrap:wrap;margin-bottom:24px;padding-bottom:16px;border-bottom:1px solid #1e1e1e}}
.mi{{font-size:12px;color:#666}}.mi strong{{color:#aaa;display:block;margin-bottom:2px}}
.cs{{margin-bottom:16px;padding:16px;background:#0d0d0d;border-radius:4px;border:1px solid #1a1a1a}}
.cs h4{{font-size:11px;color:#555;text-transform:uppercase;letter-spacing:.8px;margin-bottom:8px}}
.ct{{font-size:14px;color:#ccc;line-height:1.6}}.dt{{font-size:14px;color:#ff9800;line-height:1.6;font-style:italic}}
.nt{{font-size:12px;color:#666;margin-top:8px;padding:8px;background:#111;border-radius:3px;border-left:2px solid #333}}
.cpbtn{{background:#1e1e1e;border:1px solid #2a2a2a;color:#888;padding:6px 14px;border-radius:4px;font-size:12px;cursor:pointer;margin-top:12px}}
.cpbtn:hover{{background:#252525;color:#ccc}}
.fbtn{{background:#1a1a1a;border:1px solid #2a2a2a;color:#666;padding:5px 12px;border-radius:4px;font-size:12px;cursor:pointer}}
.fbtn:hover,.fbtn.active{{background:#222;color:#ccc;border-color:#444}}
.fbar{{display:flex;gap:8px;margin-bottom:16px;flex-wrap:wrap}}
.stitle{{font-size:13px;font-weight:600;color:#888;margin-bottom:12px;text-transform:uppercase;letter-spacing:.5px}}
.gen{{font-size:11px;color:#333;padding:16px 0 4px;text-align:right}}
.back{{color:#555;font-size:12px;cursor:pointer;margin-bottom:16px;display:inline-block}}
.back:hover{{color:#aaa}}
</style>
</head>
<body>
<div class="topbar">
  <div class="brand"><span>SumoSim</span><span class="dot">•</span>Compliance Hub</div>
  <div class="meta">Generated {generated} &nbsp;·&nbsp; {len(docs)} documents &nbsp;·&nbsp; Source: SumoSim_Compliance_Hub.xlsx</div>
</div>
<div class="layout">
  <div class="sidebar" id="sb">
    <div class="sb-sec">Overview</div>
    <div class="sb-item active" id="nav-dash" onclick="showDash(this)">📊 Dashboard</div>
    <div class="sb-sec" style="margin-top:12px">Documents</div>
  </div>
  <div class="main" id="main"></div>
</div>
<script>
const DOCS={dj},CONTENT={cj};
let CF='All';

const SC=s=>s==='Live'?'sL':s==='Draft'?'sD':s==='Under Review'?'sR':'sP';
const BC=s=>s==='Live'?'bl':s==='Draft'?'bd':s==='Under Review'?'br':'bp';
const BS=s=>s==='Live'?'L':s==='Draft'?'D':s==='Under Review'?'R':'P';
const PC=p=>p==='Yes'?'pY':p==='Pending'?'pP':'pN';
const PL=p=>p==='Yes'?'✓ Live':p==='Pending'?'⏳ Pending':'— Unpublished';

function setActive(el){{document.querySelectorAll('.sb-item').forEach(i=>i.classList.remove('active'));if(el)el.classList.add('active');}}

function buildSB(){{
  const sb=document.getElementById('sb');
  DOCS.forEach((d,i)=>{{
    const el=document.createElement('div');
    el.className='sb-item';el.id='nav-'+i;
    el.innerHTML=d.title+`<span class="badge ${{BC(d.status)}}">${{BS(d.status)}}</span>`;
    el.onclick=()=>showDoc(i,el);
    sb.appendChild(el);
  }});
}}

function showDash(el){{
  setActive(el||document.getElementById('nav-dash'));
  const cats=[...new Set(DOCS.map(d=>d.category))].filter(Boolean);
  const filt=CF==='All'?DOCS:DOCS.filter(d=>CF==='Live'?d.status==='Live':CF==='Draft'?d.status==='Draft':CF==='Review'?d.status==='Under Review':d.category===CF);
  const fbtn=['All','Live','Draft','Review',...cats].map(f=>`<button class="fbtn ${{CF===f?'active':''}}" onclick="setF('${{f}}')">${{f}}</button>`).join('');
  const live=DOCS.filter(d=>d.status==='Live').length,draft=DOCS.filter(d=>d.status==='Draft').length,pub=DOCS.filter(d=>d.published==='Yes').length;
  const rows=filt.map(d=>{{const i=DOCS.indexOf(d);return `<tr onclick="showDoc(${{i}},document.getElementById('nav-${{i}}'))">
    <td style="color:#fff;font-weight:500">${{d.title}}</td><td style="color:#666">${{d.category}}</td>
    <td style="color:#555">v${{d.version}}</td><td><span class="pill ${{SC(d.status)}}">${{d.status}}</span></td>
    <td style="color:#666">${{d.nextReview}}</td><td><span class="${{PC(d.published)}}">${{PL(d.published)}}</span></td>
  </tr>`;}}).join('');
  document.getElementById('main').innerHTML=`
    <div class="grid">
      <div class="card"><div class="num">${{DOCS.length}}</div><div class="lbl">Total Documents</div></div>
      <div class="card"><div class="num" style="color:#4caf50">${{live}}</div><div class="lbl">Live</div></div>
      <div class="card"><div class="num" style="color:#ff9800">${{draft}}</div><div class="lbl">Draft</div></div>
      <div class="card"><div class="num" style="color:#4caf50">${{pub}}</div><div class="lbl">Published on Site</div></div>
    </div>
    <div class="stitle">All Compliance Documents</div>
    <div class="fbar">${{fbtn}}</div>
    <table><thead><tr><th>Document</th><th>Category</th><th>Version</th><th>Status</th><th>Next Review</th><th>Website</th></tr></thead>
    <tbody>${{rows}}</tbody></table>
    <div class="gen">Data from SumoSim_Compliance_Hub.xlsx · Generated {generated}</div>`;
}}

function setF(f){{CF=f;showDash(document.getElementById('nav-dash'));}}

function showDoc(i,navEl){{
  setActive(navEl);
  const doc=DOCS[i],sects=CONTENT[doc.title]||[];
  const sh=sects.map(s=>`<div class="cs"><h4>${{s.section}}</h4>
    ${{s.live?`<div class="ct">${{s.live}}</div>`:''}}
    ${{s.draft&&s.draft!==s.live?`<div class="dt">📝 Draft: ${{s.draft}}</div>`:''}}
    ${{s.notes?`<div class="nt">💡 ${{s.notes}}</div>`:''}}</div>`).join('');
  const allText=sects.map(s=>s.section+'\\n'+(s.live||s.draft)).join('\\n\\n');
  document.getElementById('main').innerHTML=`
    <div class="panel">
      <span class="back" onclick="showDash(document.getElementById('nav-dash'))">← Dashboard</span>
      <div class="ptitle">${{doc.title}}</div>
      <div class="pmeta">
        <div class="mi"><strong>Category</strong>${{doc.category}}</div>
        <div class="mi"><strong>Version</strong>v${{doc.version}}</div>
        <div class="mi"><strong>Status</strong><span class="pill ${{SC(doc.status)}}" style="font-size:11px">${{doc.status}}</span></div>
        <div class="mi"><strong>Reviewed By</strong>${{doc.reviewedBy}}</div>
        <div class="mi"><strong>Last Reviewed</strong>${{doc.lastReviewed}}</div>
        <div class="mi"><strong>Next Review</strong>${{doc.nextReview}}</div>
        <div class="mi"><strong>Website Page</strong>${{doc.websitePage||'—'}}</div>
        <div class="mi"><strong>Published</strong>${{doc.published}}</div>
      </div>
      ${{sh||'<div style="color:#444;font-style:italic;padding:16px 0">No content sections yet.</div>'}}
      ${{allText.trim()?`<button class="cpbtn" onclick="navigator.clipboard.writeText(decodeURIComponent('${{encodeURIComponent(allText)}}'))
        .then(()=>this.innerText='✓ Copied!').catch(()=>this.innerText='Copy failed')">📋 Copy content for CMS</button>`:''}}
    </div>`;
}}

buildSB();
showDash(document.getElementById('nav-dash'));
</script>
</body>
</html>"""

def main():
    print("SumoSim Compliance Hub Generator v2")
    print(f"Reading {XL_FILE}...")
    docs, content = load_data(XL_FILE)
    print(f"  ✓ {len(docs)} documents, {sum(len(v) for v in content.values())} content sections")
    html = generate_html(docs, content)
    with open(OUT_FILE,"w",encoding="utf-8") as f:
        f.write(html)
    print(f"  ✓ {OUT_FILE} written ({len(html)//1024}KB)")
    print(f"\nDone! Commit index.html to GitHub → live in ~60s")
    print(f"https://donsumosec.github.io/sumosim-compliance-hub/")

if __name__=="__main__":
    main()
